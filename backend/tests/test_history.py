"""Historical workbook import: Summary is authoritative, detail only when it reconciles, no double counting."""
import shutil
from datetime import date, datetime

import pytest
from openpyxl import Workbook

from app.db import Base, SessionLocal, engine
from app.history.ingest import run_ingest
from app.models import HistIssue, HistRevenue, HistSource, Venue


def _book(path, district_detail=4000, extra_daily_row=None):
    wb = Workbook()
    s = wb.active
    s.title = "Summary"
    rows = [(None, "1 - 30 Sept"), (None, "Particulars", "Amount"), (None, None, "Pickleball"),
            (None, "Pay & Play", 7000), (None, "Online", 5000), (None, "District", 4000), (None, "Hudle", 1000),
            (None, "Offline", 2000), (None, "Bookings", 2000), (None, "Community", 500), (None, "Pickup games", 500),
            (None, "Total Gross Revenue", 7500), (None, "GST", 1144.07)]
    for r in rows:
        s.append(r)
    d = wb.create_sheet("Offline & RG Community")
    d.append(["Date", "Day", "Offline booking hours", "Booking Revenue", "Community Revenue", "Total Revenue", "Payment Gateway"])
    d.append([datetime(2026, 9, 1), "Tuesday", 1.5, 1200, 500, 1700, 30])
    d.append([datetime(2026, 9, 2), "Wednesday", 1.0, 800, 0, 800, 14])
    if extra_daily_row:
        d.append(extra_daily_row)
    d.append(["Total", None, None, 2000, 500, 2500])
    p = wb.create_sheet("District")
    p.append(["Link", "Date", "Paid to District", "Commissions (₹)", "GST on Commission", "Net Revenue", "TCS", "TDS", "Final Amount (₹)"])
    p.append([None, "1 - 30 Sept", district_detail, 280, 50, district_detail - 330, 20, 4, district_detail - 354])
    p.append([None, "Total", district_detail])
    h = wb.create_sheet("Hudle")
    h.append(["Link", "Date", "Paid to Hudle", "Commisions (Including GST)", "Net Revenue", "TDS", "TCS", "Amount"])
    h.append([None, "1 - 15 Sept", 600, 78, 522, 1, 3, 518])
    h.append([None, "16 - 30 Sept", 400, 52, 348, 1, 2, 345])
    h.append([None, None, 1000, 130, 870, None, None, 863])  # unlabelled total row
    wb.save(path)


@pytest.fixture()
def folder(tmp_path):
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    db = SessionLocal()
    db.add(Venue(code="EOK", name="East of Kailash"))
    db.commit()
    db.close()
    _book(tmp_path / "RallyGully _ EOK _ Sales Data Sept.xlsx")
    return tmp_path


def _lines(db, ch):
    return list(db.query(HistRevenue).filter(HistRevenue.channel == ch))


def test_reconciled_detail_is_used_and_totals_match(folder):
    out = run_ingest(str(folder))
    assert out["gross"] == 7500  # parent rows (Pay & Play, Online, Offline, Community) not counted
    db = SessionLocal()
    hudle = _lines(db, "hudle")
    assert [(l.date_from, l.date_to, l.gross, l.net) for l in hudle] == [
        (date(2026, 9, 1), date(2026, 9, 15), 600, 518), (date(2026, 9, 16), date(2026, 9, 30), 400, 345)]
    books = _lines(db, "offline_bookings")
    assert {(l.date_from.day, l.gross, l.sport) for l in books} == {(1, 1200, "Pickleball"), (2, 800, "Pickleball")}
    assert sum(l.hours for l in books) == 2.5
    src = db.query(HistSource).one()
    assert src.stated_total == src.ingested_total == 7500
    db.close()


def test_mismatched_detail_falls_back_to_summary(tmp_path):
    Base.metadata.drop_all(engine)
    Base.metadata.create_all(engine)
    db = SessionLocal()
    db.add(Venue(code="EOK", name="East of Kailash"))
    db.commit()
    _book(tmp_path / "RallyGully _ EOK _ Sales Data Sept.xlsx", district_detail=4500)
    out = run_ingest(str(tmp_path))
    assert out["gross"] == 7500
    d = _lines(db, "district")
    assert len(d) == 1 and d[0].gross == 4000 and d[0].basis == "month"
    assert db.query(HistIssue).filter(HistIssue.code == "detail_mismatch").count() == 1
    db.close()


def test_duplicate_copies_and_stray_dates_are_not_double_counted(folder):
    shutil.copy(folder / "RallyGully _ EOK _ Sales Data Sept.xlsx", folder / "RallyGully _ EOK _ Sales Data Sept (1).xlsx")
    out = run_ingest(str(folder))
    assert out["gross"] == 7500
    assert out["sources"] == {"ingested": 1, "duplicate": 1}
    # a stray row from another month in the daily sheet is excluded, not added
    _book(folder / "RallyGully _ EOK _ Sales Data Sept.xlsx", extra_daily_row=[datetime(2026, 8, 31), "Monday", 1, 900, 0, 900, 15])
    (folder / "RallyGully _ EOK _ Sales Data Sept (1).xlsx").unlink()
    out = run_ingest(str(folder))
    assert out["gross"] == 7500
    db = SessionLocal()
    assert all(l.date_from.month == 9 for l in db.query(HistRevenue))
    db.close()


def test_rerun_replaces_rather_than_appends(folder):
    run_ingest(str(folder))
    out = run_ingest(str(folder))
    assert out["gross"] == 7500


def test_tracker_reconciliation_sets_excluded_venues_apart(folder):
    wb = Workbook()
    s = wb.active
    s.title = "Central Dashboard"
    s.append(["Sno. ", "Venues", "April", "May", "June", "July", "August", "Sept"])
    s.append([1, "RallyGully I East of Kailash", 0, 0, 0, 0, 0, 7500.4])
    s.append([None, "Pickleball", 0, 0, 0, 0, 0, 7500])  # sport sub-row: not a venue
    s.append([2, "RallyGully I Gandhipath", 0, 0, 0, 0, 0, 900])
    s.append([None, "Total Gross Revenue", 0, 0, 0, 0, 0, 8400.4])
    wb.save(folder / "B2C Revenue Tracker _ 26-27.xlsx")
    run_ingest(str(folder))
    from app.routers.history_routes import reconciliation
    with SessionLocal() as db:
        r = reconciliation(user=None, db=db)
    assert r["summary"] == {"venue_months": 1, "matching": 0, "rounding": 1, "differs": 0, "pulse_only": 0, "tracker_only": 0}
    assert r["rows"][0]["pulse"] == 7500 and r["rows"][0]["tracker"] == 7500.4
    assert r["excluded"] == [{"label": "RallyGully I Gandhipath", "by_month": {"2026-09": 900}, "total": 900}]
    assert r["monthly"] == [{"month": "2026-09", "pulse": 7500, "excluded": 900, "tracker_total": 8400.4, "unexplained": 0.4}]


def test_monthly_revenue_uses_history_once_per_venue_month(folder):
    run_ingest(str(folder))
    from app.routers.analytics_routes import monthly
    with SessionLocal() as db:
        r = monthly(venue_id=None, user=None, db=db)
        gross = sum(l.gross for l in db.query(HistRevenue))
    assert [m["month"] for m in r["months"]] == ["2026-09"]
    assert round(r["months"][0]["revenue"], 2) == round(gross, 2) and not r["months"][0]["live"]
    assert r["venues"][0]["months"]["2026-09"] == round(gross, 2)
