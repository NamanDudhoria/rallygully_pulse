"""HQ decisions on the historical import (answers to the open questions, 14 Oct 2026).

Applied after the workbooks are imported, each one recorded as a 'hq_decision' note so the audit
trail shows why a figure differs from its source file.
"""
from sqlalchemy import select

from ..models import HistRevenue, HistSource, Venue
from .parse import CHANNELS, SPORTS, VERTICAL, norm, num
from .sources import month_bounds

# Venues that exist in the sales files but not (yet) in Pulse, created on import.
NEW_VENUES = {
    "CHD37": {"name": "RallyGully Chandigarh Sec 37", "location": "Chandigarh"},
}

DECISIONS = [
    {"kind": "move", "venue": "EOK", "from": (2026, 2), "to": (2026, 3), "channel": "memberships", "amount": 38300.0,
     "why": "HQ: the ₹38,300 Offcourt membership instalment belongs to March (file put all ₹78,300 in Feb; tracker "
            "shows ₹40,000 Feb + ₹38,300 Mar)."},
    {"kind": "add_tracker", "venue": "DW10", "month": (2026, 3), "channel": "coaching",
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 10 Dwarka", "march", "Coaching"),
     "why": "HQ: March coaching was ₹11,500 (tracker); the file's Summary shows ₹0 and its Coaching row has no amount."},
    {"kind": "use_tracker", "venue": "NOI49", "month": (2026, 4),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | Commons", "april"),
     "why": "HQ: April figures per the tracker (file and tracker disagree on offline bookings, rental, social events)."},
    {"kind": "use_tracker", "venue": "GG83", "month": (2026, 6),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | KLJ 83", "june"),
     "why": "HQ: June figures per the tracker (file and tracker disagree on offline bookings, rental, memberships)."},    {"kind": "use_tracker", "venue": "EOK", "month": (2025, 10),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully I East of Kailash", "october"),
     "why": "HQ (option A): every venue-month matches the tracker; October 2025 aligned to the tracker. Note: the tracker counts the ₹5,200 Misfits revenue twice (under Online and Community); the file's ₹1,78,232.6 adds up correctly — aligned to the tracker by HQ choice."},
    {"kind": "use_tracker", "venue": "EOK", "month": (2025, 11),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully I East of Kailash", "november"),
     "why": "HQ (option A): every venue-month matches the tracker; November 2025 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "GG83", "month": (2025, 12),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 83 Gurgaon", "december"),
     "why": "HQ (option A): every venue-month matches the tracker; December 2025 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "LCR", "month": (2026, 1),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully | LCR Pune", "january"),
     "why": "HQ (option A): every venue-month matches the tracker; January 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "DPSG", "month": (2026, 3),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully | DPSG", "march"),
     "why": "HQ (option A): every venue-month matches the tracker; March 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "DW19", "month": (2026, 3),
     "tracker": ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 19 Dwarka", "march"),
     "why": "HQ (option A): every venue-month matches the tracker; March 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "GG83", "month": (2026, 5),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | KLJ 83", "may"),
     "why": "HQ (option A): every venue-month matches the tracker; May 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "BNR", "month": (2026, 6),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | Powerplay", "june"),
     "why": "HQ (option A): every venue-month matches the tracker; June 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "DW10", "month": (2026, 6),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully I Sector 10 Dwarka", "june"),
     "why": "HQ (option A): every venue-month matches the tracker; June 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "NOI49", "month": (2026, 6),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | Commons", "june"),
     "why": "HQ (option A): every venue-month matches the tracker; June 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "DW10", "month": (2026, 7),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully I Sector 10 Dwarka", "july"),
     "why": "HQ (option A): every venue-month matches the tracker; July 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "DPSG", "month": (2026, 7),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | DPSG", "july"),
     "why": "HQ (option A): every venue-month matches the tracker; July 2026 aligned to the tracker."},
    {"kind": "use_tracker", "venue": "EOK", "month": (2026, 8),
     "tracker": ("B2C Revenue Tracker _ 26-27.xlsx", "RallyGully | EOK", "august"),
     "why": "HQ (option A): every venue-month matches the tracker; August 2026 aligned to the tracker."},
]


def ensure_venues(db) -> list[str]:
    created = []
    have = {v.code for v in db.scalars(select(Venue))}
    for code, spec in NEW_VENUES.items():
        if code not in have:
            db.add(Venue(code=code, name=spec["name"], location=spec["location"], status="active"))
            created.append(code)
    db.flush()
    return created


def _tracker_month(folder, fname, tab, col_name):
    """{channel: (total, {sport: x})} for one month column of a tracker venue tab."""
    import os
    from openpyxl import load_workbook
    ws = load_workbook(os.path.join(folder, fname), data_only=True)[tab]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    hr = next(i for i, r in enumerate(rows) if any(norm(v) in ("july", "april") for v in r)
              and any(norm(v) in ("metrics", "venues") for v in r))
    head = [norm(v) for v in rows[hr]]
    col = head.index(col_name)
    sport_cols = {}
    j = col + 1
    while j < len(head) and head[j] in SPORTS:
        sport_cols[j] = SPORTS[head[j]]
        j += 1
    lc = next(j for j, v in enumerate(head) if v in ("metrics", "venues"))
    out, total = {}, None
    for r in rows[hr + 1:]:
        lab = norm(r[lc]) if lc < len(r) else ""
        if lab.startswith("total gross"):
            total = num(r[col])
            break
        ch = CHANNELS.get(lab)
        if ch:
            sports = {sp: num(r[k]) for k, sp in sport_cols.items() if k < len(r) and num(r[k])}
            out[ch] = (num(r[col]) or 0, sports)
    return out, total


def _has_tab(folder, fname, tab) -> bool:
    import os
    from openpyxl import load_workbook
    return tab in load_workbook(os.path.join(folder, fname), read_only=True).sheetnames


def apply(run, folder) -> None:
    db = run.db
    venues = {v.code: v for v in db.scalars(select(Venue))}
    import os
    for d in DECISIONS:
        v = venues.get(d["venue"])
        tracker = d.get("tracker")
        if v is None or (tracker and not os.path.exists(os.path.join(folder, tracker[0]))):
            continue  # venue or tracker not present in this run (e.g. a partial folder)
        if tracker and not _has_tab(folder, tracker[0], tracker[1]):
            run.issue(None, "warn", "hq_decision", f"could not apply (tracker tab '{tracker[1]}' missing): {d['why']}", v.id)
            continue
        if d["kind"] == "move":
            (fy, fm), (ty, tm) = d["from"], d["to"]
            lines = list(db.scalars(select(HistRevenue).where(HistRevenue.venue_id == v.id, HistRevenue.year == fy,
                                                              HistRevenue.month == fm, HistRevenue.channel == d["channel"])))
            src_line = next((l for l in lines if l.gross >= d["amount"] - 0.5), None)
            if not src_line:
                run.issue(None, "error", "hq_decision", f"could not apply: {d['why']}", v.id, fy, fm)
                continue
            lo, hi = month_bounds(ty, tm)
            r = d["amount"] / src_line.gross
            moved = HistRevenue(**{c: getattr(src_line, c) for c in (
                "source_id", "venue_id", "vertical", "channel", "sport", "counterparty", "sheet", "row")},
                year=ty, month=tm, date_from=lo, date_to=hi, basis="month", gross=d["amount"],
                net=round(src_line.net * r, 2) if src_line.net is not None else None,
                description=f"{src_line.description + ' · ' if src_line.description else ''}moved from {fy}-{fm:02d} (HQ decision)")
            for k in ("net",):
                if getattr(src_line, k) is not None:
                    setattr(src_line, k, round(getattr(src_line, k) * (1 - r), 2))
            src_line.gross = round(src_line.gross - d["amount"], 2)
            note = f"₹{d['amount']:,.0f} moved to {ty}-{tm:02d} (HQ decision)"
            src_line.description = f"{src_line.description} · {note}" if src_line.description else note
            db.add(moved)
            run.issue(None, "info", "hq_decision", d["why"], v.id, fy, fm)
            run.issue(None, "info", "hq_decision", d["why"], v.id, ty, tm)
        elif d["kind"] == "add_tracker":
            from .ingest import _tracker_value
            y, m = d["month"]
            fname, tab, col, label = d["tracker"]
            x = _tracker_value(folder, fname, tab, col, label)
            src = HistSource(file=fname, sheet=tab, content_hash=f"decision|{fname}|{tab}|{col}|{label}", venue_id=v.id,
                             venue_label=tab, year=y, month=m, status="ingested", reason=d["why"], stated_total=x, ingested_total=x)
            db.add(src)
            db.flush()
            lo, hi = month_bounds(y, m)
            db.add(HistRevenue(source_id=src.id, venue_id=v.id, year=y, month=m, date_from=lo, date_to=hi, basis="month",
                               vertical=VERTICAL[d["channel"]], channel=d["channel"], gross=x, sheet=f"tracker › {tab}",
                               description="from B2C tracker (HQ decision)"))
            run.issue(src, "info", "hq_decision", d["why"])
        elif d["kind"] == "use_tracker":
            y, m = d["month"]
            fname, tab, col = d["tracker"]
            trk, total = _tracker_month(folder, fname, tab, col)
            src = HistSource(file=fname, sheet=tab, content_hash=f"decision|{fname}|{tab}|{col}", venue_id=v.id,
                             venue_label=tab, year=y, month=m, status="ingested",
                             reason=f"{d['why']} Tracker month total ₹{(total or 0):,.0f}; this source holds only the replaced channels.")
            db.add(src)
            db.flush()
            lines = list(db.scalars(select(HistRevenue).where(HistRevenue.venue_id == v.id, HistRevenue.year == y,
                                                              HistRevenue.month == m)))
            have = {}
            for l in lines:
                have[l.channel] = have.get(l.channel, 0) + l.gross
            lo, hi = month_bounds(y, m)
            changed, added = [], 0.0
            for ch in sorted(set(have) | set(trk)):
                t_total, t_sports = trk.get(ch, (0, {}))
                if abs(have.get(ch, 0) - t_total) <= 1:
                    continue
                for l in lines:
                    if l.channel == ch:
                        db.delete(l)
                parts = t_sports if t_sports and abs(sum(t_sports.values()) - t_total) <= 1 else {None: t_total}
                for sp, x in parts.items():
                    if x:
                        db.add(HistRevenue(source_id=src.id, venue_id=v.id, year=y, month=m, date_from=lo, date_to=hi,
                                           basis="month", vertical=VERTICAL[ch], channel=ch, sport=sp, gross=x,
                                           sheet=f"tracker › {tab}", description="from B2C tracker (HQ decision)"))
                        added += x
                changed.append(f"{ch} ₹{have.get(ch, 0):,.0f} → ₹{t_total:,.0f}")
            month_sum = sum(x for c, x in have.items() if abs(x - trk.get(c, (0, {}))[0]) <= 1) + added
            if total is not None and abs(total - month_sum) > 0.005:
                # The tab's Total Gross row disagrees with its own channel rows; HQ asked for the tracker total.
                adj = round(total - month_sum, 2)
                db.add(HistRevenue(source_id=src.id, venue_id=v.id, year=y, month=m, date_from=lo, date_to=hi,
                                   basis="month", vertical="other", channel="tracker_adjustment", gross=adj,
                                   sheet=f"tracker › {tab}",
                                   description="tracker Total Gross differs from the sum of its channel rows (HQ decision)"))
                added += adj
                changed.append(f"tracker adjustment ₹{adj:,.2f} (tracker total ₹{total:,.2f} vs channels ₹{month_sum:,.2f})")
            src.ingested_total = round(added, 2)
            run.issue(src, "info", "hq_decision", f"{d['why']} Changed: {'; '.join(changed) or 'nothing (already equal)'}")
    db.flush()
