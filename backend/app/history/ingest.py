"""Import the historical sales workbooks into the history ledger.

    python -m app.history.ingest <folder> [--report history_report.md]

Rules (no double counting, nothing invented):
  * One workbook = one venue-month. Byte-different but content-identical copies ('… (1).xlsx') are
    recorded as duplicates and skipped.
  * Each Summary's leaf rows are the authoritative monthly figure per channel; parent rows
    (Offline, Online, Community, …) are subtotals and never counted. Leaf sum must equal the
    Summary's own Total Gross Revenue.
  * A detail sheet is used for a channel only if it adds up to that Summary figure (±₹1); its
    lines then give dates, sport, fees and counterparties. Otherwise the Summary figure is stored
    as one month-level line and the mismatch is logged.
  * Months with no workbook are taken from the B2C trackers' venue tabs only where listed in
    TRACKER_GAPS, and are marked as tracker-sourced.
Re-running replaces the ledger, so it is safe to run again after corrections.
"""
import argparse
import glob
import os
from collections import defaultdict

from openpyxl import load_workbook
from sqlalchemy import delete, select

from ..db import SessionLocal, ensure_schema
from ..models import HistIssue, HistPnl, HistRevenue, HistSource, Venue
from . import parse
from .parse import VERTICAL, Line, norm
from .sources import EVIDENCE, fingerprint, month_bounds, month_from_text, venue_for_file, year_for

TOL = 1.0  # ₹ tolerance when matching detail to Summary

# Detail sheet → channels it may substantiate (first is its usual meaning).
SHEET_CHANNELS = {
    "hudle": ["hudle"], "playo": ["playo"], "district": ["district"], "khelomore": ["khelomore"],
    "misfits": ["misfits"], "courtly": ["courtly", "social_events"], "sportomic": ["sportomic"],
    "memberships": ["memberships"], "membership": ["memberships"],
    "academy": ["academy", "coaching"], "coaching": ["coaching", "academy"],
    "other income": ["other_income", "coaching", "academy"],
    "b2b": ["b2b_partnerships", "corporate_events", "brand_activation"],
    "brand activation": ["brand_activation", "b2b_partnerships", "corporate_events"],
    "corporate events": ["corporate_events", "b2b_partnerships"], "corporate event": ["corporate_events"],
    "corporateteam event": ["corporate_events"],
    "social eventstournament": ["social_events"], "social eventstournaments": ["social_events"],
    "f&b": ["fnb"],
}
DAILY_SHEETS = {"offline & rg community", "offline bookings"}

# Months with no sales workbook, taken from the trackers' venue tabs: (tracker file, tab, venue code, year, month, column header)
TRACKER_GAPS = [
    ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 10 Dwarka", "DW10", 2025, 7, "july"),
    ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 10 Dwarka", "DW10", 2025, 8, "august"),
    ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 10 Dwarka", "DW10", 2025, 9, "september"),
    ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 83 Gurgaon", "GG83", 2025, 8, "august"),
    ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 83 Gurgaon", "GG83", 2025, 9, "september"),
    ("B2C Revenue Tracker.xlsx", "RallyGully I East of Kailash", "EOK", 2025, 9, "september"),
    ("B2C Revenue Tracker _ 26-27.xlsx", "DOK | 83", "S83", 2026, 8, "august"),
]
# Summary cells that are broken (#REF!) with no detail sheet, filled from the tracker's venue tab:
# (file, channel) → (tracker file, tab, column header). Only where the tracker is the sole record.
TRACKER_FILLS = {
    ("RallyGully _ Tirang _ Sales Data Mar.xlsx", "sportomic"):
        ("B2C Revenue Tracker.xlsx", "RallyGully I Sector 10 Dwarka", "march", "Sportomic (Organized by RallyGully)"),
}
# Gandhipath (Jaipur), Ambala, Zirakpur and Golf Course EXT appear only in the trackers and are
# excluded from history by HQ decision (14 Oct 2026) — not imported, not reported.


class Run:
    def __init__(self, db):
        self.db = db
        self.venues = {v.code: v for v in db.scalars(select(Venue))}
        self.report = []

    def issue(self, src, sev, code, msg, venue_id=None, y=None, m=None):
        self.db.add(HistIssue(source_id=src.id if src else None, venue_id=venue_id or (src.venue_id if src else None),
                              year=y or (src.year if src else None), month=m or (src.month if src else None),
                              severity=sev, code=code, message=msg))

    def add_line(self, src, ln: Line, vid, y, m):
        lo, hi = month_bounds(y, m)
        f, t = ln.date_from, ln.date_to
        if t < lo or f > hi:  # dated outside its accounting month: keep the month, note the stated dates
            note = f"stated date {f}–{t}"
            ln.description = f"{ln.description} · {note}" if ln.description else note
            f, t, ln.basis = lo, hi, "month"
        else:
            f, t = max(f, lo), min(t, hi)
        self.db.add(HistRevenue(
            source_id=src.id, venue_id=vid, year=y, month=m, date_from=f, date_to=t, basis=ln.basis,
            vertical=VERTICAL[ln.channel], channel=ln.channel, sport=ln.sport, gross=round(ln.gross, 2),
            commission=ln.commission, gateway=ln.gateway, other_costs=ln.other_costs, tds=ln.tds, tcs=ln.tcs,
            net=ln.net, hours=ln.hours, people=ln.people, counterparty=(ln.counterparty or None) and ln.counterparty[:160],
            description=ln.description, sheet=ln.sheet, row=ln.row))
        return round(ln.gross, 2)


def _month_lines(ch, total, sports, single, y, m, sheet, row, why):
    lo, hi = month_bounds(y, m)
    if sports and abs(sum(sports.values()) - total) <= TOL and len(sports) > 1:
        return [Line(ch, x, sheet, row, lo, hi, "month", sport=sp, description=why) for sp, x in sports.items() if x]
    sp = single or (next(iter(sports)) if len(sports) == 1 else None)
    return [Line(ch, total, sheet, row, lo, hi, "month", sport=sp, description=why)]


def _with_sport(lines, ch, total, leaf_sports, sheet_split, single, y, m):
    """Give detail lines a sport. Lines that state one keep it; single-sport venues get theirs; on
    multi-sport venues, undated-by-sport detail is replaced by the month's sport split (which the
    Summary or the sheet states explicitly) so sport totals stay exact."""
    if all(l.sport for l in lines):
        return lines, None
    if single:
        for l in lines:
            l.sport = l.sport or single
        return lines, None
    split = None
    for cand in (leaf_sports, sheet_split):
        if cand and len(cand) > 1 and abs(sum(cand.values()) - total) <= TOL:
            split = cand
            break
    if split is None:
        return lines, None  # sport unknown — keep dated lines with sport blank
    # Month-level per-sport lines, carrying the detail's fees pro-rata.
    lo, hi = month_bounds(y, m)
    fee = {k: sum(getattr(l, k) or 0 for l in lines) for k in ("commission", "gateway", "other_costs", "tds", "tcs", "net")}
    has = {k: any(getattr(l, k) is not None for l in lines) for k in fee}
    out = []
    for sp, x in split.items():
        if not x:
            continue
        r = x / total if total else 0
        out.append(Line(ch, x, lines[0].sheet, None, lo, hi, "month", sport=sp,
                        description=f"sport split of {len(lines)} dated line(s) in '{lines[0].sheet}'",
                        **{k: round(fee[k] * r, 2) if has[k] else None for k in fee}))
    return out, f"{ch}: dated detail in '{lines[0].sheet}' has no sport; stored as the stated monthly sport split"


def ingest_file(run: Run, path: str, seen: dict):
    name = os.path.basename(path)
    db = run.db
    vm = venue_for_file(name)
    wb = load_workbook(path, data_only=True)
    if "Summary" not in wb.sheetnames:
        src = HistSource(file=name, content_hash=fingerprint(path), venue_label="?", status="excluded",
                         reason="no Summary sheet")
        db.add(src)
        db.flush()
        return
    summ = parse.parse_summary(wb["Summary"])
    m = month_from_text(name.split("Sales")[-1]) if "Sales" in name else None
    m = m or month_from_text(summ.period)
    y = year_for(m) if m else None
    code, label = vm if vm else (None, "?")
    venue = run.venues.get(code) if code else None
    h = fingerprint(path)
    src = HistSource(file=name, content_hash=h, venue_label=label, venue_id=venue.id if venue else None, year=y, month=m,
                     stated_total=summ.total, status="ingested")
    db.add(src)
    db.flush()
    if h in seen:
        src.status, src.reason = "duplicate", f"identical content to '{seen[h]}'"
        return
    seen[h] = name
    if not venue:
        src.status = "excluded"
        src.reason = "venue not in Pulse" if vm else "venue could not be identified from the file name"
        run.issue(src, "warn", "venue_missing", f"{name}: {src.reason} — not imported ({label}, total ₹{summ.total:,.0f})")
        return
    vid = venue.id
    if code in EVIDENCE and EVIDENCE[code].startswith("ASSUMED"):
        src.reason = f"venue match assumed: {EVIDENCE[code]}"
    for lab, raw in summ.errors:
        run.issue(src, "error" if parse.is_error(raw) else "warn", "summary_cell",
                  f"Summary row '{lab}' = {raw!r} (broken formula or unknown label)")
    for ch, leaf in summ.leaves.items():
        if leaf["sports"] and leaf["total"] is not None and abs(sum(leaf["sports"].values()) - leaf["total"]) > TOL:
            run.issue(src, "warn", "sport_sum", f"{leaf['label']}: sport columns add to ₹{sum(leaf['sports'].values()):,.2f} "
                      f"but its total is ₹{leaf['total']:,.2f} — sport split not used")
            leaf["sports"] = {}
    leaf_sum = sum(l["total"] or 0 for l in summ.leaves.values())
    if summ.total is not None and abs(leaf_sum - summ.total) > TOL:
        run.issue(src, "error", "summary_sum", f"leaf rows add to ₹{leaf_sum:,.2f} but Total Gross says ₹{summ.total:,.2f}")
    single = summ.single_sport or (summ.sport_cols[0] if len(summ.sport_cols) == 1 else None)

    # ── detail candidates
    daily, notes_daily = [], []
    lists = []  # (sheet title, normalized, lines, split, notes)
    for ws in wb.worksheets:
        t = norm(ws.title)
        if t in DAILY_SHEETS:
            lines, notes = parse.parse_daily(ws, y, m)
            daily.extend(lines)
            notes_daily += [f"'{ws.title}': {n}" for n in notes]
        elif t in SHEET_CHANNELS:
            lines, split, notes = parse.parse_list(ws, SHEET_CHANNELS[t][0], y, m)
            lists.append((ws.title, t, lines, split, notes))
        elif t != "summary":
            run.issue(src, "warn", "sheet_unknown", f"sheet '{ws.title}' not recognised — ignored")
    for n in notes_daily:
        if "OUTSIDE:" in n:
            run.issue(src, "warn", "daily_wrong_month", n.replace("OUTSIDE:", "") + " — this daily sheet looks copied from "
                      "another month; the Summary figures are used instead")
        else:
            run.issue(src, "info", "daily_note", n)
    for title, _, _, _, notes in lists:
        for n in notes:
            run.issue(src, "info", "sheet_note", f"'{title}': {n}")

    used_sheets = set()
    total_in = 0.0
    for ch, leaf in summ.leaves.items():
        T = leaf["total"]
        cands = []
        d = [l for l in daily if l.channel == ch]
        if d:
            cands.append(("daily", d, {}))
        for title, t, lines, split, _ in lists:
            if ch in SHEET_CHANNELS[t] and title not in used_sheets and lines:
                cands.append((title, lines, split))
        if ch == "fnb":  # a dedicated F&B sheet is richer than the daily column
            cands.sort(key=lambda c: c[0] == "daily")
        chosen = None
        if T is None:  # broken Summary cell — fall back to detail if there is any
            nz = [c for c in cands if sum(l.gross for l in c[1])]
            if len(nz) == 1:
                chosen = nz[0]
                run.issue(src, "error", "summary_broken",
                          f"{leaf['label']}: Summary cell is broken; used '{chosen[0]}' detail = ₹{sum(l.gross for l in chosen[1]):,.2f}")
            else:
                fill = TRACKER_FILLS.get((name, ch))
                x = _tracker_value(os.path.dirname(path), *fill) if fill else None
                if x is None:
                    run.issue(src, "error", "summary_broken", f"{leaf['label']}: Summary cell is broken and no detail sheet "
                              f"to recover it from — not imported")
                    continue
                run.issue(src, "warn", "summary_broken", f"{leaf['label']}: Summary cell is broken and the file has no "
                          f"detail for it; ₹{x:,.2f} taken from the tracker tab '{fill[1]}' (the only record)")
                lo, hi = month_bounds(y, m)
                total_in += run.add_line(src, Line(ch, x, f"tracker › {fill[1]}", None, lo, hi, "month", sport=single,
                                                   description="from B2C tracker: Summary cell broken in the sales file"), vid, y, m)
                continue
        else:
            for c in cands:
                if abs(sum(l.gross for l in c[1]) - T) <= TOL:
                    chosen = c
                    break
        if not T and not chosen:
            stray = [(c[0], sum(l.gross for l in c[1])) for c in cands if c[0] == "daily" or SHEET_CHANNELS[norm(c[0])][0] == ch]
            stray = [(s, x) for s, x in stray if x]
            if stray:
                run.issue(src, "warn", "detail_not_in_summary",
                          f"{ch}: Summary has ₹0 but detail shows {', '.join(f'{s} ₹{x:,.2f}' for s, x in stray)} — not imported")
            continue
        if chosen:
            if chosen[0] != "daily":
                used_sheets.add(chosen[0])
            total = T if T is not None else sum(l.gross for l in chosen[1])
            lines, note = _with_sport(chosen[1], ch, total, leaf["sports"], chosen[2], single, y, m)
            if note:
                run.issue(src, "info", "sport_split", note)
        else:
            sums = ", ".join(f"'{c[0]}' ₹{sum(l.gross for l in c[1]):,.2f}" for c in cands) or "no detail sheet"
            why = f"Summary figure (detail did not reconcile: {sums})" if cands else "Summary figure (no detail sheet)"
            if cands:
                run.issue(src, "warn", "detail_mismatch", f"{leaf['label']}: Summary ₹{T:,.2f} vs {sums} — used the Summary figure")
            lines = _month_lines(ch, T, leaf["sports"], single, y, m, "Summary", leaf["row"], why)
        for ln in lines:
            total_in += run.add_line(src, ln, vid, y, m)

    # daily columns with no Summary row at all
    for ch in sorted({l.channel for l in daily} - set(summ.leaves)):
        x = sum(l.gross for l in daily if l.channel == ch)
        if x:
            run.issue(src, "warn", "detail_not_in_summary", f"daily '{ch}' totals ₹{x:,.2f} but the Summary has no such row — not imported")
    src.ingested_total = round(total_in, 2)
    expect = summ.total if summ.total is not None else None
    if expect is not None and abs(total_in - expect) > TOL:
        run.issue(src, "error", "file_total", f"imported ₹{total_in:,.2f} ≠ Summary total ₹{expect:,.2f}")
    if summ.total is None:
        run.issue(src, "error", "file_total", f"Summary Total Gross is broken; imported ₹{total_in:,.2f} from leaf rows and detail")

    for lab, amt, note in summ.pnl:
        l = lab.lower()
        if amt is None and not note:
            continue
        cat = ("gst" if "gst" in l and "post gst" not in l else
               "commissions" if "commission" in l and "net" not in l and "post commission" not in l else
               "net" if l.startswith(("total net", "net revenue", "total revenue", "realised")) else
               "expenses_total" if l.startswith("total expenses") or l == "total" else
               "profit" if l.startswith(("net profit", "net loss", "final profit")) else
               "share" if "share" in l else
               "other" if amt is None else "expense")
        db.add(HistPnl(source_id=src.id, venue_id=vid, year=y, month=m, category=cat, label=lab[:200], amount=amt, note=note))


def _tracker_value(folder, fname, tab, col_name, row_label):
    ws = load_workbook(os.path.join(folder, fname), data_only=True)[tab]
    rows = [[c.value for c in r] for r in ws.iter_rows()]
    hr = next(i for i, r in enumerate(rows) if any(norm(v) in ("july", "april") for v in r))
    col = next(j for j, v in enumerate(rows[hr]) if norm(v) == col_name)
    for r in rows[hr + 1:]:
        if any(norm(v) == norm(row_label) for v in r[:3]):
            return parse.num(r[col])
    return None


def ingest_tracker_gaps(run: Run, folder: str):
    db = run.db
    for fname, tab, code, y, m, col_name in TRACKER_GAPS:
        path = os.path.join(folder, fname)
        venue = run.venues.get(code)
        if not os.path.exists(path) or not venue:
            continue
        ws = load_workbook(path, data_only=True)[tab]
        rows = [[c.value for c in r] for r in ws.iter_rows()]
        hr = next(i for i, r in enumerate(rows) if any(norm(v) in ("july", "april") for v in r))
        col = next(j for j, v in enumerate(rows[hr]) if norm(v) == col_name)
        lc = next(j for j, v in enumerate(rows[hr]) if norm(v) in ("metrics", "venues"))
        src = HistSource(file=fname, sheet=tab, content_hash=f"{fname}|{tab}|{col_name}", venue_id=venue.id,
                         venue_label=tab, year=y, month=m, status="ingested",
                         reason="no sales workbook for this month — figures from the B2C tracker's venue tab")
        db.add(src)
        db.flush()
        total = None
        tot_in = 0.0
        lo, hi = month_bounds(y, m)
        for i, r in enumerate(rows[hr + 1:], start=hr + 2):
            lab = norm(r[lc]) if lc < len(r) else ""
            if not lab:
                continue
            v = parse.num(r[col]) if col < len(r) else None
            if lab.startswith("total gross"):
                total = v
                break
            ch = parse.CHANNELS.get(lab)
            if ch and v:
                tot_in += run.add_line(src, Line(ch, v, tab, i, lo, hi, "month", description="from B2C tracker (no sales file)"),
                                       venue.id, y, m)
        src.stated_total, src.ingested_total = total, round(tot_in, 2)
        sev = "warn" if total is None or abs(total - tot_in) > TOL else "info"
        run.issue(src, sev, "tracker_month", f"{tab} {m}/{y}: no sales workbook; imported ₹{tot_in:,.2f} from the tracker "
                  f"(tracker total ₹{(total or 0):,.2f}). Trackers were found to differ from workbooks in other months — verify.")


def run_ingest(folder: str, report_path: str | None = None) -> dict:
    ensure_schema()
    db = SessionLocal()
    try:
        for model in (HistIssue, HistPnl, HistRevenue, HistSource):
            db.execute(delete(model))
        from . import decisions
        created = decisions.ensure_venues(db)
        run = Run(db)
        for code in created:
            run.issue(None, "info", "venue_created", f"Created venue {code} ({decisions.NEW_VENUES[code]['name']}) — "
                      f"add its courts, hours and prices under Configuration")
        seen: dict = {}
        files = sorted(f for f in glob.glob(os.path.join(folder, "*.xlsx")) if not os.path.basename(f).startswith(("B2C", "~$")))
        # originals before their '(1)' copies so the copy is the one marked duplicate
        files.sort(key=lambda f: ("(1)" in f, f))
        for f in files:
            ingest_file(run, f, seen)
        ingest_tracker_gaps(run, folder)
        decisions.apply(run, folder)
        db.commit()
        out = summarize(db)
        if report_path:
            write_report(db, report_path, out)
        return out
    finally:
        db.close()


def summarize(db) -> dict:
    srcs = list(db.scalars(select(HistSource)))
    by = defaultdict(int)
    for s in srcs:
        by[s.status] += 1
    lines = list(db.scalars(select(HistRevenue)))
    issues = defaultdict(int)
    for i in db.scalars(select(HistIssue)):
        issues[i.severity] += 1
    return {"sources": dict(by), "lines": len(lines), "gross": round(sum(l.gross for l in lines), 2),
            "issues": dict(issues)}


def write_report(db, path, out):
    venues = {v.id: v for v in db.scalars(select(Venue))}
    srcs = list(db.scalars(select(HistSource).order_by(HistSource.venue_id, HistSource.year, HistSource.month)))
    L = ["# Historical sales import — report", "",
         f"Sources: {out['sources']} · revenue lines: {out['lines']} · gross imported: ₹{out['gross']:,.2f} · issues: {out['issues']}", "",
         "## Venue matching", ""]
    for code, ev in EVIDENCE.items():
        L.append(f"- **{code}**: {ev}")
    L += ["", "## Sources", "", "| Venue | Month | File | Status | Summary total | Imported | Note |", "|---|---|---|---|---:|---:|---|"]
    for s in srcs:
        v = venues.get(s.venue_id)
        L.append(f"| {v.code if v else s.venue_label} | {s.year or ''}-{(s.month or 0):02d} | {s.file}{' › ' + s.sheet if s.sheet else ''} | "
                 f"{s.status} | {'' if s.stated_total is None else f'{s.stated_total:,.2f}'} | "
                 f"{'' if s.ingested_total is None else f'{s.ingested_total:,.2f}'} | {(s.reason or '')[:120]} |")
    L += ["", "## Issues", ""]
    for sev in ("error", "warn", "info"):
        items = list(db.scalars(select(HistIssue).where(HistIssue.severity == sev)))
        L.append(f"### {sev} ({len(items)})")
        src = {s.id: s for s in srcs}
        for i in items:
            s = src.get(i.source_id)
            L.append(f"- [{i.code}] {s.file if s else ''}: {i.message}")
        L.append("")
    with open(path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(L))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("folder")
    ap.add_argument("--report", default="history_report.md")
    a = ap.parse_args()
    out = run_ingest(a.folder, a.report)
    print(out)
    print(f"Report written to {a.report}")


if __name__ == "__main__":
    main()
