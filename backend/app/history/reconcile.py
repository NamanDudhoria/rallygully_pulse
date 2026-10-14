"""Capture the B2C trackers' Central Dashboard (venue × month and Total Gross) for reconciliation."""
import os

from openpyxl import load_workbook
from sqlalchemy import select

from ..models import HistTracker, Venue
from .parse import SPORTS, norm
from .sources import MONTHS

# Dashboard row → Pulse venue code. Rows not listed (Gandhipath, Ambala, Zirakpur, Golf Course EXT.) are venues
# HQ excluded from Pulse; they are kept with no venue so the tracker's grand total can still be explained.
ROWS = [("sector 10 dwarka", "DW10"), ("tirang", "DW10"), ("sector 83 gurgaon", "GG83"), ("blue court", "GG83"),
        ("east of kailash", "EOK"), ("drive or kick | sec 83", "S83"), ("sector 19 dwarka", "DW19"),
        ("drive or kick | dwarka", "DW19"), ("dpsg", "DPSG"), ("panchsheel", "PSP"), ("lcr pune", "LCR"),
        ("law college", "LCR"), ("powerplay", "BNR"), ("commons", "NOI49"), ("chandigarh", "CHD37")]
TRACKERS = {"B2C Revenue Tracker.xlsx": 2025, "B2C Revenue Tracker _ 26-27.xlsx": 2026}  # file → FY start year


def code_for(label: str) -> str | None:
    lab = norm(label)
    return next((c for k, c in ROWS if k in lab), None)


def capture(db, folder: str) -> int:
    codes = {v.code: v.id for v in db.scalars(select(Venue))}
    n = 0
    for fname, fy in TRACKERS.items():
        path = os.path.join(folder, fname)
        if not os.path.exists(path):
            continue
        wb = load_workbook(path, data_only=True)
        if "Central Dashboard" not in wb.sheetnames:
            continue
        rows = [[c.value for c in r] for r in wb["Central Dashboard"].iter_rows()]
        hr = next((i for i, r in enumerate(rows) if any(norm(v) == "venues" for v in r)), None)
        if hr is None:
            continue
        head = rows[hr]
        lc = next(j for j, v in enumerate(head) if norm(v) == "venues")
        cols = {}
        for j, h in enumerate(head):
            t = norm(h)
            if "(" in t:  # e.g. "April (26-27)" — next year's column, captured from its own tracker
                continue
            m = MONTHS.get(t.split(" ")[0]) if t else None
            if m:
                cols[j] = (fy if m >= 4 else fy + 1, m)
        for r in rows[hr + 1:]:
            label = str(r[lc] or "").strip() if lc < len(r) else ""
            lab = norm(label)
            if not lab or lab in SPORTS:
                continue  # blank or a sport sub-row under a venue
            kind = "total" if lab.startswith("total gross") else "venue"
            code = code_for(label) if kind == "venue" else None
            for j, (y, m) in cols.items():
                x = r[j] if j < len(r) else None
                if isinstance(x, (int, float)) and x:
                    db.add(HistTracker(file=fname, label=label, venue_id=codes.get(code), kind=kind, year=y, month=m,
                                       amount=float(x)))
                    n += 1
            if kind == "total":
                break
    return n
