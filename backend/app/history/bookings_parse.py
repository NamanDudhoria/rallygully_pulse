"""Read the platform booking exports (Hudle, Playo, Khelomore, District) into one shape.

Each export lists bookings with the slot played, court, sport, customer name and money. One parsed
record is one booking on one court for one continuous time range. Rows that are totals, blank,
cancelled or refunds are dropped here and counted, so the import report can say what was skipped.
"""
import csv
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from .parse import norm, num

# Sheet or venue name → Pulse venue code. None = a venue HQ excluded (not imported). Order matters:
# the Sec 83 Drive or Kick names must be matched before the Dwarka Drive or Kick ones.
VENUE_PATTERNS = [
    (r"jaipur|ambala|zirakpur|gandhipath|golf course", None),
    (r"(drive or kick|dok)\W*(\||_)?\W*sec\W*83|83\W*(drive or kick|dok)", "S83"),
    (r"drive (?:and|or) kick|19 dwarka", "DW19"),
    (r"tirang|dwarka 10|10 dwarka", "DW10"),
    (r"blue court|83 gurgaon|sec 83 gurgaon|^sec 83$|sector 83", "GG83"),
    (r"east of kailash|(?<![a-z])eok(?![a-z])", "EOK"),
    (r"law college|lcr", "LCR"),
    (r"dpsg|palam vihar", "DPSG"),
    (r"panchsheel", "PSP"),
    (r"commons|sector 49 noida", "NOI49"),
    (r"powerplay", "BNR"),
    (r"chandigarh", "CHD37"),
]

PLATFORMS = ("hudle", "playo", "khelomore", "district")


def venue_code(*names) -> tuple[bool, str | None]:
    """(matched, code). matched with code None means an excluded venue."""
    for name in names:
        low = norm(name).replace("⁠", "")
        if not low:
            continue
        for pat, code in VENUE_PATTERNS:
            if re.search(pat, low):
                return True, code
    return False, None


def sport_of(text) -> str | None:
    t = norm(text)
    if not t or t == "-":
        return None
    if "pickle" in t:
        return "Pickleball"
    if "padel" in t:
        return "Padel"
    if "cricket" in t:
        return "Cricket"
    if "football" in t or "turf" in t or "7v7" in t or "5v5" in t:
        return "Turf"
    if "basket" in t:
        return "Basketball"
    return None


def name_key(name) -> str:
    """Normalised customer name for matching: lower case, letters and spaces only."""
    s = re.sub(r"[^a-z ]+", " ", norm(name))
    return re.sub(r"\s+", " ", s).strip()


@dataclass
class Booking:
    platform: str
    venue: str
    day: date
    start_min: int
    end_min: int
    court: str
    sport: str | None
    amount: float  # booking value of this court-range in rupees (gross, before platform commission)
    medium: str  # online | offline
    customer: str
    ref: str
    file: str
    sheet: str
    row: int
    parts: dict = field(default_factory=dict)  # money columns as exported, for choosing what counts as revenue


@dataclass
class ParseResult:
    bookings: list = field(default_factory=list)
    skipped: dict = field(default_factory=dict)  # reason → count
    unknown_venues: dict = field(default_factory=dict)  # label → rows

    def skip(self, reason, n=1):
        self.skipped[reason] = self.skipped.get(reason, 0) + n


def _clock(text) -> int | None:
    """'05:30 PM' / '5:30pm' / '17:30:00' → minutes after midnight."""
    if text is None:
        return None
    if isinstance(text, datetime):
        return text.hour * 60 + text.minute
    if hasattr(text, "hour"):
        return text.hour * 60 + text.minute
    t = str(text).strip().lower().replace(".", "")
    m = re.match(r"(\d{1,2}):(\d{2})(?::\d{2})?\s*(am|pm)?$", t)
    if not m:
        m = re.match(r"(\d{1,2})\s*(am|pm)$", t)
        if not m:
            return None
        h, mi, ap = int(m.group(1)), 0, m.group(2)
    else:
        h, mi, ap = int(m.group(1)), int(m.group(2)), m.group(3)
    if ap == "pm" and h != 12:
        h += 12
    if ap == "am" and h == 12:
        h = 0
    return h * 60 + mi


def _ranges(pairs: list[tuple[int, int]]) -> list[tuple[int, int]]:
    """Merge back-to-back ranges: [(17:00,17:30),(17:30,18:00)] → [(17:00,18:00)]."""
    out = []
    for a, b in sorted(pairs):
        if b <= a:
            b += 24 * 60  # runs past midnight
        if out and a == out[-1][1]:
            out[-1] = (out[-1][0], b)
        else:
            out.append((a, b))
    return out


def _date(v, dayfirst=False) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = str(v or "").strip()
    for fmt in (("%d/%m/%Y", "%m/%d/%Y") if dayfirst else ("%m/%d/%Y", "%d/%m/%Y")) + ("%Y-%m-%d", "%d-%m-%Y", "%b %d, %Y"):
        try:
            return datetime.strptime(s[:10] if fmt == "%Y-%m-%d" else s, fmt).date()
        except ValueError:
            continue
    return None


def _rows(path):
    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        yield ws.title, [list(r) for r in ws.iter_rows(values_only=True)]


def _header(rows, must):
    for i, r in enumerate(rows[:6]):
        cells = [norm(c) for c in r]
        if must in cells:
            return i, cells
    return None, None


# ── Hudle ────────────────────────────────────────────────────────────────────
HUDLE_STATUS = {"confirmed", "cancelled", "rescheduled", "partially_cancelled", "completed", "pending", "failed"}
HUDLE_PAY = ("pay at venue", "hudle wallet", "venue wallet", "hudle passes", "qr payment", "online paid amount")
_HUD_RANGE = re.compile(r"(\d{1,2}:\d{2}\s*[AP]M)\s*-\s*(\d{1,2}:\d{2}\s*[AP]M)", re.I)
_HUD_DATE = re.compile(r"([A-Z][a-z]{2} \d{1,2}, \d{4})")
_COURT = re.compile(r"(court|turf|7v7|5v5|box|pitch|padel|arena)\b", re.I)


def _hud_days(text) -> dict:
    """'Oct 10, 2025 05:30 PM - 06:00 PM | Oct 11, 2025 05:30 PM - 06:00 PM' → {day: [(start, end), …]}."""
    out: dict = {}
    for seg in str(text or "").split("|"):
        dm = _HUD_DATE.search(seg)
        if dm and _date(dm.group(1)):
            out.setdefault(_date(dm.group(1)), []).extend((_clock(a), _clock(b)) for a, b in _HUD_RANGE.findall(seg))
    return out


def parse_hudle(path, res: ParseResult):
    fname = os.path.basename(path)
    for sheet, rows in _rows(path):
        hi, head = _header(rows, "slot details")
        if hi is None:
            continue
        col = {h: j for j, h in enumerate(head) if h}
        for i, r in enumerate(rows[hi + 1:], start=hi + 2):
            cells = {h: (r[j] if j < len(r) else None) for h, j in col.items()}
            ref = next((str(v) for v in r if isinstance(v, str) and re.fullmatch(r"HUD\d+", v.strip())), None)
            if not ref and "booking reference id" not in col and cells.get("venue") and norm(r[0]) != "total":
                # Some sheets (Dec 2025 EOK, DPSG) were exported without the reference column: key the booking
                # on its payment id, or on its place in the file.
                ref = str(cells.get("payment gateway ids") or "").strip() or f"{fname}|{sheet}|{i}"
                anchor = col["source"]  # where the reference column sits in the usual layout
                r = r[:anchor] + [ref] + r[anchor:]
            if not ref:
                if any(v not in (None, "") for v in r):
                    res.skip("hudle: total or unlabelled row")
                continue
            # Entry type is found by content too: in shifted rows its column holds an amount or a date.
            entry = next((norm(v) for v in r if isinstance(v, str) and norm(v).endswith("_entry")), "")
            if entry not in ("", "original_entry"):
                res.skip("hudle: adjustment entry (not a booking)")
                continue
            # Exports mix layouts: some rows are shifted (a missing 'Cancelled Slot Details' cell) and carry
            # extra columns after the payment block. Payment cells are read by name, realigned on the booking
            # reference; status, sport and court are found by their content.
            shift = r.index(ref) - col.get("booking reference id", r.index(ref) - 1)
            g = lambda h: r[col[h] + shift] if h in col and 0 <= col[h] + shift < len(r) else None  # noqa: E731
            statuses = [norm(v) for v in r if isinstance(v, str) and norm(v) in HUDLE_STATUS]
            status = statuses[0] if statuses else "confirmed"
            if status not in ("confirmed", "partially_cancelled", "completed"):
                res.skip(f"hudle: status {status}")
                continue
            ok, code = venue_code(g("venue"), sheet)
            if not ok:
                res.unknown_venues[f"hudle | {sheet} | {g('venue')}"] = res.unknown_venues.get(f"hudle | {sheet} | {g('venue')}", 0) + 1
                continue
            if code is None:
                res.skip("excluded venue (Jaipur, Ambala, …)")
                continue
            slot_cells = [v for v in r[:r.index(ref) + 2] if isinstance(v, str) and _HUD_RANGE.search(v)]
            days = _hud_days(slot_cells[0] if slot_cells else "")
            for d, cs in _hud_days(slot_cells[1] if len(slot_cells) > 1 else "").items():
                for c in cs:
                    if c in days.get(d, []):
                        days[d].remove(c)
            days = {d: ps for d, ps in days.items() if ps}
            if not days or any(p[0] is None or p[1] is None for ps in days.values() for p in ps):
                res.skip("hudle: unreadable slot details")
                continue
            parts = {h: num(g(h)) or 0 for h in HUDLE_PAY}
            amount = sum(parts.values())
            booked = next((_date(v) for v in r[:r.index(ref)] if isinstance(v, (datetime, date)) or
                           (isinstance(v, str) and re.fullmatch(r"\d{1,4}[-/]\d{1,2}[-/]\d{2,4}", v.strip()))), None)
            court_cell = next((v for v in r if isinstance(v, str) and _COURT.match(v.strip())), None)
            courts = [c.strip() for c in str(court_cell or "").split(",") if c.strip()] or ["?"]
            sport = next((sport_of(v) for v in r[r.index(ref) + 1:] if isinstance(v, str) and not v.startswith("http")
                          and sport_of(v)), None)
            medium = "offline" if norm(g("source")) == "offline" else "online"  # entered by the venue, paid there
            # A slot listed twice on a day means two courts at that time; give each court its own ranges. More
            # copies than named courts: the extra copies were on courts the export didn't name.
            spans = []
            for day, pairs in sorted(days.items()):
                per_court: dict[str, list] = {c: [] for c in courts}
                seen: dict = {}
                for p in pairs:
                    k = seen.get(p, 0)
                    seen[p] = k + 1
                    per_court.setdefault(courts[k] if k < len(courts) else f"{courts[-1]} +{k - len(courts) + 1}",
                                         []).append(p)
                spans += [(day, c, a, b) for c, ps in per_court.items() for a, b in _ranges(ps)]
            total_min = sum(b - a for _, _, a, b in spans) or 1
            for day, c, a, b in spans:
                share = (b - a) / total_min
                res.bookings.append(Booking("hudle", code, day, a, b, c, sport,
                                            amount * share, medium, str(r[0] or "").strip(),
                                            ref, fname, sheet, i, {**{k: v * share for k, v in parts.items()}, "booked_on": booked}))


# ── Playo ────────────────────────────────────────────────────────────────────
PLAYO_ALIASES = {"bookingid": "id", "slot_date": "date", "sport_name": "sport", "user_name": "name"}  # Powerplay layout
_PLAYO_RANGE = re.compile(r"\s*(\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AP]M)?)\s*-\s*(\d{1,2}:\d{2}(?::\d{2})?\s*(?:[AP]M)?)\s*$", re.I)


def parse_playo(path, res: ParseResult):
    fname = os.path.basename(path)
    for sheet, rows in _rows(path):
        hi, head = _header(rows, "gross")
        if hi is None:
            continue
        col = {PLAYO_ALIASES.get(h, h): j for j, h in enumerate(head) if h}
        price = "net" if "net" in col else "gross"  # that layout lists the venue's discount; net is what was charged
        groups: dict = {}
        for i, r in enumerate(rows[hi + 1:], start=hi + 2):
            g = lambda h: r[col[h]] if h in col and col[h] < len(r) else None  # noqa: E731
            bid = str(g("id") or "").strip()
            if bid.lower() in ("id", "bookingid"):
                continue  # the header repeated above a second table
            if not bid or norm(g("timestamp")).startswith("total") or (g("date") is None and not g("time")):
                if any(v not in (None, "") for v in r):
                    res.skip("playo: total or unlabelled row")
                continue
            ok, code = venue_code(sheet)
            if not ok:
                res.unknown_venues[f"playo | {sheet}"] = res.unknown_venues.get(f"playo | {sheet}", 0) + 1
                continue
            if code is None:
                res.skip("excluded venue (Jaipur, Ambala, …)")
                continue
            day = _date(g("date"))  # Playo writes the slot date month-first (04/07/2026 = 7 April)
            t = str(g("time") or "")
            m = _PLAYO_RANGE.match(t)
            if not day or not m:
                res.skip("playo: unreadable date or time")
                continue
            court = str(g("court") or "?").strip().rstrip(",") or "?"
            key = (bid, day, court)
            groups.setdefault(key, {"pairs": [], "amount": 0.0, "sport": g("sport"), "name": g("name"), "row": i,
                                    "code": code, "booked": _date(str(g("timestamp") or "")[:10], dayfirst=True)})
            groups[key]["pairs"].append((_clock(m.group(1)), _clock(m.group(2))))
            groups[key]["amount"] += num(g(price)) or 0
        for (bid, day, court), x in groups.items():
            spans = _ranges(x["pairs"])
            total_min = sum(b - a for a, b in spans) or 1
            for a, b in spans:
                res.bookings.append(Booking("playo", x["code"], day, a, b, court, sport_of(x["sport"]),
                                            x["amount"] * (b - a) / total_min, "online", str(x["name"] or "").strip(),
                                            bid, fname, sheet, x["row"], {"booked_on": x["booked"]}))


# ── Khelomore ────────────────────────────────────────────────────────────────
KHELOMORE_ALIASES = {"booking id": "bookingid", "slot date": "slot_date", "start time": "slot_start_time",
                     "end time": "slot_end_time", "total actual price": "slot_amount", "venue": "venuename",
                     "property": "propertyname"}  # the older export layout (to Nov 2025)


def parse_khelomore(path, res: ParseResult):
    fname = os.path.basename(path)
    for sheet, rows in _rows(path):
        hi, head = _header(rows, "bookingid")
        if hi is None:
            hi, head = _header(rows, "booking id")
            if hi is None:
                continue
            head = [KHELOMORE_ALIASES.get(h, h) for h in head]
        col = {h: j for j, h in enumerate(head) if h}
        for i, r in enumerate(rows[hi + 1:], start=hi + 2):
            g = lambda h: r[col[h]] if h in col and col[h] < len(r) else None  # noqa: E731
            bid = str(g("bookingid") or "").strip()
            if bid.lower() in ("bookingid", "booking id"):
                continue  # the header repeated above a second table
            if not bid or bid.lower().startswith("total"):
                if any(v not in (None, "") for v in r):
                    res.skip("khelomore: total or unlabelled row")
                continue
            ok, code = venue_code(g("venuename"), sheet)
            if not ok:
                res.unknown_venues[f"khelomore | {sheet} | {g('venuename')}"] = 1
                continue
            if code is None:
                res.skip("excluded venue (Jaipur, Ambala, …)")
                continue
            day, a, b = _date(g("slot_date")), _clock(g("slot_start_time")), _clock(g("slot_end_time"))
            if not day or a is None or b is None:
                res.skip("khelomore: unreadable date or time")
                continue
            if b <= a:
                b += 24 * 60
            amount = (num(g("slot_amount")) or 0) - (num(g("discountedslotamount")) or 0) * 0  # list price paid
            res.bookings.append(Booking("khelomore", code, day, a, b, str(g("propertyname") or "?").strip(),
                                        sport_of(g("propertyname")) or sport_of(sheet), amount, "online",
                                        str(g("username") or "").strip(), bid, fname, sheet, i))


# ── District ─────────────────────────────────────────────────────────────────
def parse_district_xlsx(path, res: ParseResult, year: int):
    """Monthly District payout sheet: one row per 30-minute slot sold ('January 10 | 8:00am')."""
    fname = os.path.basename(path)
    for sheet, rows in _rows(path):
        hi, head = _header(rows, "slot details")
        if hi is None:
            continue
        col = {h: j for j, h in enumerate(head) if h}
        for i, r in enumerate(rows[hi + 1:], start=hi + 2):
            g = lambda h: r[col[h]] if h in col and col[h] < len(r) else None  # noqa: E731
            slot = str(g("slot details") or "")
            m = re.match(r"\s*([A-Za-z]+)\s+(\d{1,2})\s*\|\s*(\d{1,2}:\d{2}\s*[ap]m)", slot, re.I)
            if not m:
                if any(v not in (None, "") for v in r):
                    res.skip("district: total or unlabelled row")
                continue
            ok, code = venue_code(g("facility name"), sheet)
            if not ok or code is None:
                res.skip("excluded venue (Jaipur, Ambala, …)") if ok else res.unknown_venues.update({f"district | {sheet}": 1})
                continue
            day = _date(f"{m.group(1)[:3]} {int(m.group(2))}, {year}")
            a = _clock(m.group(3))
            res.bookings.append(Booking("district", code, day, a, a + 30, "?", None, num(g("ticket sales")) or 0,
                                        "online", "", f"{sheet}:{i}", fname, sheet, i))


def parse_district_csv(path, res: ParseResult):
    fname = os.path.basename(path)
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    refunded: dict = {}
    for r in rows:  # refund rows reference the sale they reverse (same Order ID)
        if r.get("Action") == "refund":
            refunded[r["Order ID"]] = refunded.get(r["Order ID"], 0) + (num(r.get("Refund Amount")) or num(r.get("Total Amount")) or 0)
    for i, r in enumerate(rows, start=2):
        if not r.get("Order ID"):
            res.skip("district: blank row")
            continue
        if r.get("Action") != "sale":
            res.skip("district: refund row (netted against its sale)")
            continue
        status = r.get("Order Status")
        if status in ("refund", "cancelled"):
            res.skip(f"district: order {status}")
            continue
        ok, code = venue_code(r.get("Venue"))
        if not ok:
            res.unknown_venues[f"district | {r.get('Venue')}"] = res.unknown_venues.get(f"district | {r.get('Venue')}", 0) + 1
            continue
        if code is None:
            res.skip("excluded venue (Jaipur, Ambala, …)")
            continue
        day, a, b = _date(r.get("Slot Start Date")), _clock(r.get("Slot Start Time")), _clock(r.get("Slot End Time"))
        if not day or a is None or b is None:
            res.skip("district: unreadable date or time")
            continue
        if b <= a:
            b += 24 * 60
        amount = (num(r.get("Total Amount")) or 0) - refunded.get(r["Order ID"], 0)
        if amount <= 0:
            res.skip("district: fully refunded")
            continue
        medium = "offline" if norm(r.get("Booking Medium")) == "offline" else "online"
        parts = {k: num(r.get(k)) or 0 for k in ("Slot Amount", "Add on Amount", "Total Amount", "Received Amount",
                                                   "Remaining Amount", "Discount", "Commissionable Amount")}
        parts["refunded"] = refunded.get(r["Order ID"], 0)
        parts["booked_on"] = _date(r.get("Booking Date"))
        res.bookings.append(Booking("district", code, day, a, b, (r.get("Arena / Court") or "?").strip(),
                                    sport_of(r.get("Sport")), amount, medium, (r.get("Customer Name") or "").strip(),
                                    r["Order ID"], fname, "csv", i, parts))


def parse_folder(folder: str) -> tuple[ParseResult, list[dict]]:
    """Parse every export in ``folder``. Byte-identical or content-identical copies are read once."""
    import hashlib
    res = ParseResult()
    files = []
    seen: dict = {}
    for f in sorted(os.listdir(folder), key=lambda f: ("(" in f, f)):
        p = os.path.join(folder, f)
        if not os.path.isfile(p) or f.startswith("~$"):
            continue
        low = f.lower()
        platform = ("district" if "district" in low else "hudle" if "hudle" in low else "playo" if "playo" in low
                    else "khelomore" if "khelomore" in low else None)
        entry = {"file": f, "platform": platform, "status": "read", "bookings": 0}
        files.append(entry)
        if platform is None:
            entry["status"] = "skipped: not a booking export"
            continue
        before = len(res.bookings)
        skipped_before = dict(res.skipped)
        try:
            if f.endswith(".csv"):
                parse_district_csv(p, res)
            elif platform == "district":
                parse_district_xlsx(p, res, year=2026 if re.search(r"jan|feb|mar", low) else 2025)
            elif platform == "hudle":
                parse_hudle(p, res)
            elif platform == "playo":
                parse_playo(p, res)
            else:
                parse_khelomore(p, res)
        except Exception as e:  # noqa: BLE001 — one broken export must not stop the rest
            entry["status"] = f"error: {e}"
            continue
        new = res.bookings[before:]
        digest = hashlib.sha1(repr(sorted((b.platform, b.venue, b.day, b.start_min, b.end_min, b.court, round(b.amount, 2),
                                           b.ref) for b in new)).encode()).hexdigest()
        if digest in seen:
            del res.bookings[before:]
            res.skipped = skipped_before  # a copy's skipped rows were already counted with the original
            entry["status"] = f"duplicate of {seen[digest]}"
            continue
        seen[digest] = f
        entry["bookings"] = len(new)
    # The same booking can sit in two exports (re-downloaded months, overlapping date ranges): keep it once.
    unique, dupes = {}, {}
    for b in res.bookings:
        k = (b.platform, b.ref, b.day, b.start_min, b.court)
        if k in unique:
            dupes[b.file] = dupes.get(b.file, 0) + 1
            continue
        unique[k] = b
    res.bookings = list(unique.values())
    for e in files:
        if dupes.get(e["file"]):
            e["status"] = f"read ({dupes[e['file']]} bookings already in another export)"
            e["bookings"] -= dupes[e["file"]]
    if dupes:
        res.skip("booking already in another export", sum(dupes.values()))
    return res, files
