"""Parse player lists pasted from a sheet / WhatsApp, or uploaded as .xlsx / .csv.

Two shapes are understood:

* Tables (pasted from Excel / Sheets, or uploaded): a header row (Name, Phone, Level)
  or none, in which case columns are recognised by content.
* The Community team's WhatsApp sign-up post: a preamble, then sections such as
  "7-8 PM | Intermediate" with numbered names ("1. Jatan", "2. Jatan + 1", "4." left
  empty). Phones are optional; "+ 1" adds a guest. Section time and level carry over
  to the names under them.

Output rows keep their source line so the preview can point at what needs fixing.
"""
import csv
import io
import re

from ..audit import DomainError
from .customers import normalize_phone

MAX_ROWS = 500
MAX_BYTES = 2 * 1024 * 1024

LEVEL_WORDS = {
    "beginner": "beginner", "beg": "beginner", "b": "beginner", "novice": "beginner", "basic": "beginner",
    "intermediate": "intermediate", "inter": "intermediate", "int": "intermediate", "i": "intermediate",
    "mid": "intermediate", "medium": "intermediate",
    "advanced": "advanced", "adv": "advanced", "a": "advanced", "pro": "advanced", "expert": "advanced",
}
HEADERS = {
    "name": ("name", "player", "full name", "player name", "customer"),
    "phone": ("phone", "mobile", "number", "contact", "whatsapp", "phone number", "mobile number", "ph", "mob"),
    "level": ("level", "skill", "skill level", "category", "grade"),
}
_PHONE_RE = re.compile(r"(\+?\d[\d\s\-()]{8,16}\d)")


def level_of(value: str) -> str | None:
    return LEVEL_WORDS.get((value or "").strip().lower().rstrip("."))


def _phone_or_none(value: str) -> str | None:
    try:
        return normalize_phone(value)
    except DomainError:
        return None


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # Excel stores phone numbers as floats
    return str(v).strip()


def rows_from_text(text: str) -> list[list[str]]:
    out = []
    for line in (text or "").splitlines():
        if not line.strip():
            out.append([])
            continue
        if "\t" in line:
            cells = line.split("\t")
        elif "," in line or ";" in line:
            cells = next(csv.reader([line], delimiter=";" if line.count(";") > line.count(",") else ","))
        else:
            # Free text such as "Meera Kapoor 98666 66666 advanced" or "+91 98666 66666 - Meera"
            m = _PHONE_RE.search(line)
            if m:
                rest = (line[:m.start()] + " " + line[m.end():]).strip(" -:|·")
                words = rest.split()
                lvl = [w for w in words if level_of(w) and len(w) > 1]
                name = " ".join(w for w in words if w not in lvl).strip(" -:|·")
                cells = [name, m.group(1)] + lvl[:1]
            else:
                cells = [line]
        out.append([c.strip() for c in cells])
    return out


def rows_from_file(filename: str, data: bytes) -> list[list[str]]:
    if len(data) > MAX_BYTES:
        raise DomainError("File is larger than 2 MB.", 422)
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 — any parse failure is a user-facing format error
            raise DomainError("Couldn't read that Excel file. Save it as .xlsx or .csv and try again.", 422) from None
        ws = wb.worksheets[0]
        rows = [[_cell(v) for v in r] for r, _ in zip(ws.iter_rows(values_only=True), range(MAX_ROWS + 50))]
        wb.close()
        return rows
    if name.endswith((".csv", ".txt")):
        text = data.decode("utf-8-sig", errors="replace")
        return [[c.strip() for c in r] for r in csv.reader(io.StringIO(text))]
    if name.endswith(".xls"):
        raise DomainError("Old .xls files aren't supported — save as .xlsx or .csv.", 422)
    raise DomainError("Upload an .xlsx or .csv file.", 422)


def _header_map(row: list[str]) -> dict[str, int] | None:
    if any(_phone_or_none(c) for c in row):
        return None
    found = {}
    for i, c in enumerate(row):
        key = c.strip().lower()
        for field, names in HEADERS.items():
            if field not in found and key in names:
                found[field] = i
    return found if "phone" in found else None


def parse(rows: list[list[str]]) -> list[dict]:
    """→ [{line, name, phone, level, error}] with blank rows dropped."""
    indexed = [(i + 1, r) for i, r in enumerate(rows) if any(c.strip() for c in r)]
    if not indexed:
        return []
    header = _header_map(indexed[0][1])
    if header:
        indexed = indexed[1:]
    if len(indexed) > MAX_ROWS:
        raise DomainError(f"Up to {MAX_ROWS} players at a time.", 422)
    out = []
    for line, r in indexed:
        name = phone = level = None
        if header:
            get = lambda f: r[header[f]].strip() if f in header and header[f] < len(r) else ""  # noqa: E731
            phone, name = _phone_or_none(get("phone")), get("name") or None
            level = level_of(get("level"))
            raw_phone = get("phone")
        else:
            raw_phone = ""
            rest = []
            for c in r:
                if not c:
                    continue
                if phone is None and _phone_or_none(c):
                    phone, raw_phone = _phone_or_none(c), c
                elif level is None and level_of(c):
                    level = level_of(c)
                else:
                    rest.append(c)
            name = next((c for c in rest if not re.fullmatch(r"[\d\s+\-().]+", c)), None)
            if phone is None:
                raw_phone = next((c for c in rest if re.search(r"\d{5,}", c)), "")
        if not phone and not raw_phone and not name:
            continue
        error = None if phone or not raw_phone else f"'{raw_phone}' isn't a valid mobile"  # no phone: name only
        out.append({"line": line, "name": name, "phone": phone or raw_phone or None, "level": level, "guests": 0,
                    "section": None, "error": error})
    return out


# ── WhatsApp sign-up posts ───────────────────────────────────────────────────
_TIME = r"(\d{1,2})(?::(\d{2}))?\s*(am|pm)?"
_SECTION_RE = re.compile(rf"^\W*{_TIME}\s*(?:-|–|—|to)\s*{_TIME}\b\s*(.*)$", re.I)
_NUMBERED_RE = re.compile(r"^\s*(?:\d{1,3}\s*[.)\-:]|[-•*·▪●])\s*(.*)$")
_GUESTS_RE = re.compile(r"^(.*?)\s*\+\s*(\d{1,2})\s*$")


def _to_min(h: str, m: str | None, ampm: str | None) -> int:
    hour = int(h) % 12 if ampm else int(h)
    if ampm and ampm.lower() == "pm":
        hour += 12
    return hour * 60 + int(m or 0)


def section_of(line: str) -> dict | None:
    """"7-8 PM | Intermediate" → {label, start_min, end_min, level}."""
    m = _SECTION_RE.match(line.strip())
    if not m or _PHONE_RE.search(line):
        return None
    h1, m1, ap1, h2, m2, ap2, rest = m.groups()
    if int(h1) > 23 or int(h2) > 23:
        return None
    ap1 = ap1 or ap2
    start, end = _to_min(h1, m1, ap1), _to_min(h2, m2, ap2)
    if ap2 and not m.group(3) and start > end:  # "11-1 PM" → 11 AM–1 PM
        start -= 12 * 60
    words = re.split(r"[|/,·\-–]+|\s+", rest or "")
    level = next((level_of(w) for w in words if level_of(w) and len(w) > 1), None)
    return {"label": line.strip().strip("*_ "), "start_min": start, "end_min": end, "level": level}


def _entry(text: str) -> tuple[str | None, str | None, str | None, int, str | None]:
    """Free text for one player → (name, phone, level, guests, error)."""
    text = text.strip().strip("*_~")
    cells = [c.strip() for c in re.split(r"\t|,|;", text) if c.strip()] if re.search(r"[\t,;]", text) else [text]
    phone = raw_phone = level = None
    rest = []
    for c in cells:
        m = _PHONE_RE.search(c) if phone is None else None
        if m and _phone_or_none(m.group(1)):
            phone, raw_phone = _phone_or_none(m.group(1)), m.group(1)
            c = (c[:m.start()] + " " + c[m.end():]).strip(" -:|·")
        if len(cells) > 1 and level is None and level_of(c):
            level = level_of(c)
            continue
        if c:
            rest.append(c)
    name = " ".join(rest).strip(" -:|·")
    words = name.split()
    if level is None and len(words) > 1 and level_of(words[-1]) and len(words[-1]) > 1:
        level, name = level_of(words[-1]), " ".join(words[:-1])
    guests = 0
    g = _GUESTS_RE.match(name)
    if g:
        name, guests = g.group(1).strip(), int(g.group(2))
    error = None
    if phone is None:
        digits = re.search(r"\d[\d\s\-()]{3,}\d", name)
        if digits:
            raw_phone = digits.group(0)
            name = (name[:digits.start()] + name[digits.end():]).strip(" -:|·")
            error = f"'{raw_phone.strip()}' isn't a valid mobile"
    return name or None, phone or (raw_phone.strip() if raw_phone else None), level, guests, error


def parse_post(text: str) -> tuple[list[dict], list[dict]]:
    """WhatsApp post or pasted list → (rows, sections)."""
    rows: list[dict] = []
    sections: list[dict] = []
    current = None
    started = False
    header = None
    for i, raw in enumerate((text or "").splitlines(), start=1):
        line = raw.strip()
        if not line:
            continue
        sec = section_of(line)
        if sec:
            current = {**sec, "count": 0}
            sections.append(current)
            started = True
            continue
        num = _NUMBERED_RE.match(line)
        if num:
            content = num.group(1).strip()
            if not content:
                continue  # empty slot "4."
            started = True
        else:
            content = line
            if "\t" in line and header is None and not rows:
                cells = [c.strip() for c in line.split("\t")]
                header = _header_map(cells)
                if header:
                    started = True
                    continue
            if header:
                cells = [c.strip() for c in line.split("\t")]
                get = lambda f: cells[header[f]] if f in header and header[f] < len(cells) else ""  # noqa: E731
                content = "\t".join(x for x in (get("name"), get("phone"), get("level")) if x)
            elif not started and not _PHONE_RE.search(line):
                continue  # preamble: "OpenPlay Today", "📍 venue", "Please list your names below:"
            elif not _PHONE_RE.search(line) and (len(line.split()) > 5 or line.endswith(":")):
                continue  # chatter between sections
        started = True
        name, phone, level, guests, error = _entry(content)
        if not name and not phone:
            continue
        row = {"line": i, "name": name, "phone": phone, "level": level or (current or {}).get("level"),
               "guests": guests, "section": current["label"] if current else None, "error": error}
        if current:
            current["count"] += 1 + guests
        rows.append(row)
    if len(rows) > MAX_ROWS:
        raise DomainError(f"Up to {MAX_ROWS} players at a time.", 422)
    return rows, sections
